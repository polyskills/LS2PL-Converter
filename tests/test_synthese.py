"""
Tests de core.lightspeed_synthese : consolidation des exports Lightspeed
« Tickets » + « Transactions » en classeur de synthèse par période de service.

Les exports sont fabriqués ici de toutes pièces (comme _build_sample_xlsx dans
test_pipeline.py) plutôt que repris d'un fichier réel : le dépôt est public,
un export Lightspeed contient le détail des ventes d'un client identifiable.

Lancer avec : python -m pytest tests/ -q
"""
import csv
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import pytest
from openpyxl import Workbook

from core.lightspeed_synthese import (
    SITES,
    SyntheseError,
    classer_fichiers,
    construire_synthese,
    deviner_site,
    famille,
)

# Colonnes réellement exploitées par le traitement, dans l'ordre du vrai export.
COLS_TICKETS = ["Identifier", "Date", "Account", "AccountName", "Total", "PreTax",
                "Couverts", "Type", "Annulée", "Profil", "OpenDate"]
COLS_TRANSACTIONS = ["Identifier", "Account", "Type", "Qty", "UnitPrice", "FinalPrice",
                     "SKU", "Item", "Group", "TaxName", "TaxRate", "PreTax", "TaxAmount"]

# Un ticket par situation à couvrir : période normale, période d'ouverture qui
# diffère de celle du règlement, ticket d'après minuit rattaché à la veille,
# et annulation. Les totaux tickets sont exactement la somme des lignes de
# transaction correspondantes — c'est le contrôle central de l'outil.
TICKETS = [
    # (Identifier, Date règlement, Account, AccountName, Total, PreTax, Couverts, Type, Annulée, Profil Lightspeed, OpenDate)
    ("R1", "07/09/26 13:30", "A1", "BAR, Table 5", 100, 90, 2, "SALE", "Non", "Bar Journée", "07/09/26 12:00"),
    ("R2", "07/09/26 20:15", "A2", "BAR, Table 9", 60, 50, 3, "SALE", "Non", "Bar Soir", "07/09/26 17:30"),
    ("R3", "08/09/26 02:30", "A3", "BAR, Table 3", 40, 35, 1, "SALE", "Non", "Bar Nuit", "08/09/26 02:00"),
    ("R4", "07/09/26 13:10", "A4", "BAR, Table 5", -20, -18, -1, "VOID", "Oui", "Bar Journée", "07/09/26 13:00"),
]

TRANSACTIONS = [
    # (Identifier, Account, Type, Qty, FinalPrice, Item, Group, PreTax, TaxAmount)
    ("S1", "A1", "SALE", 1, 60, "Planche", "FOOD Snack", 54, 6),
    ("S2", "A1", "SALE", 2, 40, "Bière", "BEV Biere", 36, 4),
    ("S3", "A2", "SALE", 3, 60, "Cocktail maison", "Cocktail", 50, 10),
    ("S4", "A3", "SALE", 1, 40, "Plat du soir", "MP CUISINE", 35, 5),
    ("S5", "A4", "VOID", -1, -20, "Planche", "FOOD Snack", -18, -2),
]


def _classeur(entetes: list[str], lignes: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(entetes)
    for l in lignes:
        ws.append(l)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _exports(tickets=TICKETS, transactions=TRANSACTIONS):
    """(liste tickets, liste transactions) au format attendu par
    construire_synthese : des couples (nom de fichier, octets). Extension
    .xlsx — un vrai export Lightspeed est en .xls, mais l'écrire demanderait
    une dépendance d'écriture supplémentaire (xlwt) alors que seule la
    LECTURE du .xls est nécessaire en production (xlrd, déjà présent)."""
    lt = [[i, d, a, an, tot, pre, cv, ty, ann, pr, od]
          for i, d, a, an, tot, pre, cv, ty, ann, pr, od in tickets]
    lx = [[i, a, ty, q, fp / q if q else fp, fp, "SKU", it, gr, "TVA", 0.1, pre, tax]
          for i, a, ty, q, fp, it, gr, pre, tax in transactions]
    return (
        [("client_bar_tickets_20260907.xlsx", _classeur(COLS_TICKETS, lt))],
        [("client_bar_transactions_20260907.xlsx", _classeur(COLS_TRANSACTIONS, lx))],
    )


def _contenu(classeur: bytes) -> dict:
    """Valeurs de toutes les cellules, onglet par onglet — de quoi comparer
    deux classeurs sans dépendre de leur horodatage interne."""
    wb = openpyxl.load_workbook(io.BytesIO(classeur))
    return {s: [[c.value for c in r] for r in wb[s].iter_rows()] for s in wb.sheetnames}


def _csv(entetes: list[str], lignes: list[list], sep=";", decimal=",") -> bytes:
    """Export CSV à la française : séparateur point-virgule et virgule
    décimale, encodé en cp1252 — la combinaison que produit un Lightspeed
    configuré en locale française, et celle qui casse une lecture naïve."""
    def cellule(v):
        if isinstance(v, float):
            return str(v).replace(".", decimal)
        return "" if v is None else str(v)

    # Écrit avec le module csv, donc avec les guillemets là où il en faut :
    # AccountName vaut "BAR, Table 5" et casserait un fichier à séparateur
    # virgule si on se contentait de concaténer.
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=sep, lineterminator="\r\n")
    writer.writerow(entetes)
    for l in lignes:
        writer.writerow([cellule(v) for v in l])
    return buf.getvalue().encode("cp1252")


def _exports_csv():
    """Mêmes données que _exports(), au format CSV."""
    lt = [[i, d, a, an, float(tot), float(pre), cv, ty, ann, pr, od]
          for i, d, a, an, tot, pre, cv, ty, ann, pr, od in TICKETS]
    lx = [[i, a, ty, q, float(fp) / q if q else float(fp), float(fp), "SKU", it, gr, "TVA", 0.1, float(pre), float(tax)]
          for i, a, ty, q, fp, it, gr, pre, tax in TRANSACTIONS]
    return (
        [("client_bar_tickets_20260907.csv", _csv(COLS_TICKETS, lt))],
        [("client_bar_transactions_20260907.csv", _csv(COLS_TRANSACTIONS, lx))],
    )


def test_famille_prefixe_puis_table_explicite():
    assert famille("FOOD Snack") == "FOOD"
    assert famille("BEV Biere") == "BEV"
    assert famille("DIV Divers") == "DIV"
    assert famille("Cocktail") == "BEV"          # sans préfixe : table explicite
    assert famille("Cuisine Froid") == "FOOD"
    assert famille("MP CUISINE") == "AUTRE"      # inconnu : jamais absorbé en silence


def test_classer_fichiers_repartit_par_convention_de_nommage():
    t, x, inconnus = classer_fichiers([
        "cli_bar_tickets_20260907_20260908.xls",
        "cli_bar_transactions_20260907_20260908.xls",
        "export_retravaille.xls",
    ])
    assert t == ["cli_bar_tickets_20260907_20260908.xls"]
    assert x == ["cli_bar_transactions_20260907_20260908.xls"]
    assert inconnus == ["export_retravaille.xls"]


def test_synthese_totaux_et_controle_dequilibre():
    tickets, transactions = _exports()
    res = construire_synthese(tickets, transactions, "BAR")

    assert res.nb_tickets == 4
    assert res.nb_lignes == 5
    # CA = somme de TOUTES les lignes, VOID compris (ils se compensent)
    assert res.ca_ttc == 180.0
    assert res.ca_ht == 157.0
    # Les couverts d'un ticket annulé sont négatifs et annulent ceux du ticket repris
    assert res.couverts == 5
    # Contrôle central : total transactions - total tickets doit être nul
    assert res.ecart_controle == 0.0
    assert res.sans_anomalie_bloquante


def test_synthese_periode_a_louverture_et_journee_dexploitation():
    tickets, transactions = _exports()
    res = construire_synthese(tickets, transactions, "BAR")
    wb = openpyxl.load_workbook(io.BytesIO(res.classeur))
    lignes = list(wb["TICKETS"].iter_rows(values_only=True))
    entetes = list(lignes[0])
    par_id = {l[entetes.index("Identifier")]: l for l in lignes[1:]}
    periode = lambda ident: par_id[ident][entetes.index("Profil")]

    assert periode("R1") == "Bar Journée"
    # Ouvert à 17h30 (Afterwork) mais réglé à 20h15 : c'est l'OUVERTURE qui compte
    assert periode("R2") == "Bar Afterwork"
    assert par_id["R2"][entetes.index("ProfilReglement")] == "Bar Soir"
    # Ouvert à 2h du matin le 08 : rattaché à la journée d'exploitation du 07
    assert periode("R3") == "Bar Nuit"
    assert res.jours == [__import__("datetime").date(2026, 9, 7)]
    assert res.periode_libelle == "07/09/2026"


def test_synthese_signale_les_groupes_non_mappes_et_les_annulations():
    tickets, transactions = _exports()
    res = construire_synthese(tickets, transactions, "BAR")
    libelles = {a[0]: a for a in res.anomalies}

    assert libelles["Groupes non mappés (famille AUTRE)"][1] == 1
    assert "MP CUISINE" in libelles["Groupes non mappés (famille AUTRE)"][2]
    assert libelles["Tickets annulés (VOID, comptés en négatif)"][1] == 1
    # R2 : profil Lightspeed (clôture) "Bar Soir" vs période retenue (ouverture)
    cle = "Tickets rattachés à leur période d'ouverture, et non au profil Lightspeed de clôture"
    assert libelles[cle][1] == 1
    assert "R2" in libelles[cle][2]


def test_synthese_produit_les_sept_onglets_attendus():
    tickets, transactions = _exports()
    res = construire_synthese(tickets, transactions, "BAR")
    wb = openpyxl.load_workbook(io.BytesIO(res.classeur))
    assert wb.sheetnames == ["SYNTHESE", "JOUR x PERIODE", "ROTATIONS",
                             "DUREE PRESENCE", "DONNEES", "TICKETS", "ANOMALIES"]


def test_synthese_site_restaurant_a_ses_propres_periodes():
    tickets, transactions = _exports()
    res = construire_synthese(tickets, transactions, "RESTAURANT")
    wb = openpyxl.load_workbook(io.BytesIO(res.classeur))
    colonne_a = [r[0] for r in wb["SYNTHESE"].iter_rows(values_only=True)]
    assert "Restaurant Midi" in colonne_a and "Restaurant Soir" in colonne_a
    assert "Bar Journée" not in colonne_a
    assert [p for p, *_ in SITES["RESTAURANT"]["periodes"]] == ["Restaurant Midi", "Restaurant Soir"]


def test_synthese_le_site_est_un_parametre_pas_un_etat_global():
    # Deux consolidations de sites différents à la suite ne doivent pas
    # s'influencer : le script d'origine mutait un global au lancement, ce qui
    # aurait mélangé deux traitements dans le process partagé de Streamlit.
    tickets, transactions = _exports()
    bar1 = construire_synthese(tickets, transactions, "BAR")
    construire_synthese(tickets, transactions, "RESTAURANT")
    bar2 = construire_synthese(tickets, transactions, "BAR")
    # Comparaison du contenu et non des octets : un .xlsx embarque son
    # horodatage de création, deux classeurs identiques n'ont donc jamais
    # exactement les mêmes octets.
    assert _contenu(bar1.classeur) == _contenu(bar2.classeur)


def test_synthese_rapports_intervertis_message_explicite():
    tickets, transactions = _exports()
    with pytest.raises(SyntheseError) as e:
        construire_synthese(transactions, tickets, "BAR")
    assert "intervertis" in str(e.value)


def test_synthese_site_inconnu_refuse():
    tickets, transactions = _exports()
    with pytest.raises(SyntheseError, match="inconnu"):
        construire_synthese(tickets, transactions, "BRASSERIE")


def test_synthese_sans_fichier_refuse():
    _, transactions = _exports()
    with pytest.raises(SyntheseError, match="Aucun fichier"):
        construire_synthese([], transactions, "BAR")


def test_deviner_site_depuis_le_nom_des_exports():
    # Le nom porte la source, mais pas forcément le mot "bar" : ici la marque
    # de l'établissement ("barutopic") suffit.
    assert deviner_site(["anne-sophiepic-paris_barutopic_tickets_20260907.xls"]) == "BAR"
    assert deviner_site(["client_restaurant_transactions_20260907.xls"]) == "RESTAURANT"


def test_deviner_site_ne_tranche_pas_quand_cest_ambigu():
    # Deux sites évoqués, ou aucun : mieux vaut ne rien proposer qu'imposer un
    # site arbitraire — c'est le choix à l'écran qui détermine le calcul.
    assert deviner_site(["bar_tickets.xls", "restaurant_tickets.xls"]) is None
    assert deviner_site(["export_20260907.xls"]) is None
    assert deviner_site([]) is None


def test_synthese_lit_aussi_les_exports_csv():
    # Même contenu en .xlsx et en .csv français (point-virgule, virgule
    # décimale, cp1252) : les deux doivent donner exactement le même classeur.
    res_xlsx = construire_synthese(*_exports(), "BAR")
    res_csv = construire_synthese(*_exports_csv(), "BAR")

    assert res_csv.ca_ttc == res_xlsx.ca_ttc == 180.0
    assert res_csv.ca_ht == res_xlsx.ca_ht == 157.0
    assert res_csv.couverts == res_xlsx.couverts == 5
    assert res_csv.ecart_controle == 0.0
    assert _contenu(res_csv.classeur) == _contenu(res_xlsx.classeur)


def test_synthese_csv_a_separateur_virgule():
    # Séparateur détecté, pas supposé : un export en locale anglaise passe aussi.
    lt = [[i, d, a, an, tot, pre, cv, ty, ann, pr, od]
          for i, d, a, an, tot, pre, cv, ty, ann, pr, od in TICKETS]
    lx = [[i, a, ty, q, fp / q if q else fp, fp, "SKU", it, gr, "TVA", 0.1, pre, tax]
          for i, a, ty, q, fp, it, gr, pre, tax in TRANSACTIONS]
    tickets = [("bar_tickets.csv", _csv(COLS_TICKETS, lt, sep=",", decimal="."))]
    transactions = [("bar_transactions.csv", _csv(COLS_TRANSACTIONS, lx, sep=",", decimal="."))]

    res = construire_synthese(tickets, transactions, "BAR")
    assert res.ca_ttc == 180.0
    assert res.ecart_controle == 0.0


def test_synthese_melange_xlsx_et_csv():
    # Rien n'impose que les deux rapports aient le même format.
    tickets_xlsx, _ = _exports()
    _, transactions_csv = _exports_csv()
    res = construire_synthese(tickets_xlsx, transactions_csv, "BAR")
    assert res.ca_ttc == 180.0
    assert res.ecart_controle == 0.0


def test_le_rattachement_des_periodes_nest_pas_un_point_a_verifier():
    # Une table ouverte avant une frontière de période et réglée après en
    # déclenche une : dans un bar, tous les jours. C'est la règle de calcul de
    # l'outil, pas une anomalie — elle figure dans le classeur pour expliquer
    # un écart avec un rapport Lightspeed natif, mais ne doit ni alerter, ni
    # faire passer la consolidation en avertissement.
    res = construire_synthese(*_exports(), "BAR")
    libelles_verif = [a[0] for a in res.anomalies_a_verifier]
    assert not any(l.startswith("Tickets rattachés") for l in libelles_verif)
    assert not any(l.startswith("Écart total transactions") for l in libelles_verif)
    # Mais la ligne reste présente dans le classeur
    assert any(a[0].startswith("Tickets rattachés") for a in res.anomalies)


def test_journee_sans_rattachement_ni_groupe_inconnu_na_rien_a_verifier():
    # Jeu réduit : un seul ticket, ouvert et réglé dans la même période, sur un
    # groupe mappé. Aucun point à vérifier, donc statut OK côté historique.
    tickets = [("R1", "07/09/26 13:30", "A1", "BAR, Table 5", 100, 90, 2, "SALE", "Non",
                "Bar Journée", "07/09/26 12:00")]
    transactions = [("S1", "A1", "SALE", 1, 100, "Planche", "FOOD Snack", 90, 10)]
    res = construire_synthese(*_exports(tickets, transactions), "BAR")
    assert res.anomalies_a_verifier == []
    assert res.sans_anomalie_bloquante


# --- Journée sans vente ----------------------------------------------------
#
# Cas réel : restaurant fermé. LightSpeed produit quand même ses deux
# rapports, vides. La consolidation le refusait, et cinq journées sont restées
# bloquées douze jours. Même décision que pour la conversion comptable, où une
# journée sans vente a cessé d'être une erreur.


def _exports_vides(periode="20260912_20260913"):
    return (
        [(f"cli_bar_tickets_{periode}.xlsx", _classeur(COLS_TICKETS, []))],
        [(f"cli_bar_transactions_{periode}.xlsx", _classeur(COLS_TRANSACTIONS, []))],
    )


def test_journee_sans_vente_produit_un_classeur_a_zero():
    tickets, transactions = _exports_vides()
    res = construire_synthese(tickets, transactions, "BAR", periode=("12/09/26", "13/09/26"))

    assert res.sans_vente
    assert res.ca_ttc == 0.0 and res.ca_ht == 0.0 and res.couverts == 0
    assert res.nb_tickets == 0 and res.nb_lignes == 0
    assert res.sans_anomalie_bloquante          # rien de cassé : c'est une journée fermée
    assert res.anomalies_a_verifier == []       # ni erreur, ni point à vérifier
    wb = openpyxl.load_workbook(io.BytesIO(res.classeur))
    assert wb.sheetnames[0] == "SYNTHESE"


def test_journee_sans_vente_datee_depuis_le_nom_de_fichier():
    # Le contenu ne peut plus dire de quelle journée il s'agit : la période
    # vient du nom de l'export.
    tickets, transactions = _exports_vides()
    res = construire_synthese(tickets, transactions, "BAR", periode=("12/09/26", "13/09/26"))
    assert res.periode_libelle == "12/09/2026"
    entete = openpyxl.load_workbook(io.BytesIO(res.classeur))["SYNTHESE"]["A1"].value
    assert "12/09/2026" in entete


def test_journee_sans_vente_est_signalee_explicitement():
    # Jamais silencieuse : la ligne doit figurer dans le classeur pour qu'on
    # sache, en relisant, que la journée était vide et non ratée.
    tickets, transactions = _exports_vides()
    res = construire_synthese(tickets, transactions, "BAR", periode=("12/09/26", "13/09/26"))
    libelles = [a[0] for a in res.anomalies]
    assert "Aucune vente sur la période" in libelles
    detail = next(a[2] for a in res.anomalies if a[0] == "Aucune vente sur la période")
    assert "export LightSpeed qu'il faut vérifier" in detail

    anomalies_feuille = list(openpyxl.load_workbook(io.BytesIO(res.classeur))["ANOMALIES"]
                             .iter_rows(values_only=True))
    assert any(r[0] == "Aucune vente sur la période" for r in anomalies_feuille)


def test_journee_sans_vente_sans_periode_ne_plante_pas():
    # Nom de fichier sans période lisible : on préfère un classeur daté du jour
    # à un plantage ou à un refus de traiter.
    tickets, transactions = _exports_vides()
    res = construire_synthese(tickets, transactions, "BAR")
    assert res.sans_vente and res.jours and res.classeur


def test_une_journee_avec_ventes_nest_pas_marquee_sans_vente():
    res = construire_synthese(*_exports(), "BAR")
    assert not res.sans_vente
