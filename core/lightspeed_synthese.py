"""
Synthèse du CA par période de service, à partir des exports Lightspeed Back
Office « Tickets » et « Transactions » (.xls) — indépendante de la conversion
comptable vers Pennylane : elle ne consulte pas le référentiel du client et ne
produit pas d'écriture, seulement un classeur d'analyse d'exploitation.

Portage du script `synthese_lightspeed.py` validé en ligne de commande. La
logique métier est reprise telle quelle ; seules trois choses changent, parce
qu'un service web n'est pas une commande shell :

- les périodes de service ne sont plus un global muté au lancement mais un
  paramètre explicite : le process Streamlit est partagé, deux consolidations
  BAR et RESTAURANT lancées de près se mélangeraient sinon ;
- les erreurs lèvent SyntheseError au lieu d'appeler sys.exit() ;
- tout circule en mémoire (octets en entrée, octets en sortie), comme
  build_pennylane_csv, plutôt que par des chemins de fichiers.

Règles métier (inchangées) :
- Le CA d'un ticket = somme de TOUTES ses lignes de transaction, tous types
  confondus (SALE, SPLIT, UPDATE, TRANSFER, VOID, RECALL, FOREIGN). Les types
  "techniques" se compensent ; leur somme reconstitue exactement le total ticket.
- La période retenue est celle de l'OUVERTURE de la table, pas du règlement :
  une commande passée en Afterwork et encaissée en Soir compte en Afterwork.
  La déclinaison au règlement figure en second bloc de la synthèse.
- Les couverts viennent des tickets (les VOID sont négatifs et annulent).
- Famille = préfixe du groupe Lightspeed (BEV / FOOD / DIV), avec une table de
  correspondance explicite pour les groupes hors nomenclature.
- Journée d'exploitation : les tickets après minuit et avant 05h30 sont
  rattachés à la veille.
"""
from __future__ import annotations

import csv
import io
import re
import datetime as dt
from dataclasses import dataclass, field

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class SyntheseError(Exception):
    """Fichier illisible, manquant, ou dont la structure n'est pas celle
    attendue d'un export Lightspeed. Équivalent de LightspeedParseError pour
    la conversion comptable : bloque avec un message affichable tel quel."""


# (nom, libellé, heure de début en minutes depuis minuit). La période d'un
# ticket = celle dans laquelle tombe l'OUVERTURE de la table (OpenDate).
# `titre` est l'intitulé porté par les blocs de la feuille SYNTHESE.
SITES = {
    "BAR": {
        "titre": "UTOPIC",
        "periodes": [
            ("Bar Journée", "(5H30-16H59)", 5 * 60 + 30),
            ("Bar Afterwork", "(17H00-18H59)", 17 * 60),
            ("Bar Soir", "(19H00-21H59)", 19 * 60),
            ("Bar Nuit", "(22H00-5H29)", 22 * 60),
        ],
    },
    "RESTAURANT": {
        "titre": "ASPP",
        "periodes": [
            ("Restaurant Midi", "(5H30-16H59)", 5 * 60 + 30),
            ("Restaurant Soir", "(17H00-5H29)", 17 * 60),
        ],
    },
}

DEBUT_JOURNEE = pd.Timedelta(hours=5, minutes=30)

# Correspondance explicite groupe -> famille pour les groupes sans préfixe
# normalisé. Un groupe qui n'entre dans aucun cas tombe en "AUTRE" et est
# signalé dans l'onglet ANOMALIES, jamais absorbé silencieusement.
MAPPING_GROUPES = {
    "Cocktail": "BEV",
    "Cuisine chaud": "FOOD",
    "Cuisine Chaud": "FOOD",
    "Cuisine Froid": "FOOD",
    "Patisserie": "FOOD",
}

TAUPE, GRIS_CLAIR, GRIS_MOYEN, ANTHRACITE = "B7B09C", "B5B7BB", "898989", "333333"
FONT = "Montserrat"

# Lignes de l'onglet ANOMALIES qui ne signalent PAS un problème et ne doivent
# donc ni faire passer une consolidation en avertissement, ni s'afficher comme
# un point à vérifier :
# - le contrôle d'équilibre, toujours présent et déjà remonté en indicateur ;
# - le rattachement des tickets à leur période d'OUVERTURE, qui se déclenche
#   pour toute table ouverte avant une frontière de période et réglée après.
#   Autant dire tous les jours dans un bar (9 tickets sur 23 sur une journée
#   réelle) : le présenter comme une anomalie ferait croire à un problème
#   récurrent alors que c'est la règle de calcul assumée de l'outil. La ligne
#   reste dans le classeur, où elle sert à expliquer un écart avec un rapport
#   Lightspeed natif, mais comme information.
# L'ancien libellé est conservé pour que les consolidations déjà archivées
# s'affichent de la même façon que les nouvelles.
LIBELLES_INFORMATIFS = (
    "Aucune vente sur la période",
    "Écart total transactions",
    "Tickets rattachés à leur période d'ouverture",
    "Tickets dont la période (ouverture) diffère du profil Lightspeed",
)


def est_informatif(libelle: str) -> bool:
    return str(libelle).startswith(LIBELLES_INFORMATIFS)

# Colonnes indispensables au traitement : contrôlées à la lecture pour
# transformer un KeyError pandas illisible en message actionnable, et pour
# détecter tout de suite qu'on a interverti les deux rapports.
COLONNES_TICKETS = ["Identifier", "Date", "OpenDate", "Account", "AccountName", "Total",
                    "PreTax", "Couverts", "Type", "Annulée", "Profil"]
COLONNES_TRANSACTIONS = ["Identifier", "Account", "Type", "Qty", "FinalPrice", "PreTax",
                         "TaxAmount", "TaxName", "Item", "Group"]

# Colonnes sur lesquelles portent les calculs. Lues depuis un .xls/.xlsx elles
# arrivent déjà typées ; depuis un .csv elles arrivent en texte, avec une
# virgule décimale si l'export a été fait en locale française — d'où la
# conversion explicite, sans laquelle les totaux tomberaient tous à zéro.
COLONNES_NUMERIQUES_TICKETS = ["Total", "PreTax", "Couverts"]
COLONNES_NUMERIQUES_TRANSACTIONS = ["Qty", "FinalPrice", "PreTax", "TaxAmount"]


@dataclass
class SyntheseResult:
    """Résultat d'une consolidation : le classeur généré et de quoi le
    présenter à l'écran comme dans l'historique, sans le rouvrir."""
    site: str
    classeur: bytes = b""
    fichiers_sources: list[str] = field(default_factory=list)
    jours: list = field(default_factory=list)
    nb_tickets: int = 0
    nb_lignes: int = 0
    ca_ttc: float = 0.0
    ca_ht: float = 0.0
    couverts: int = 0
    anomalies: list = field(default_factory=list)

    @property
    def periode_libelle(self) -> str:
        if not self.jours:
            return ""
        if len(self.jours) == 1:
            return f"{self.jours[0]:%d/%m/%Y}"
        return f"{self.jours[0]:%d/%m/%Y} au {self.jours[-1]:%d/%m/%Y}"

    @property
    def ecart_controle(self) -> float:
        """Écart total transactions - tickets (TTC) relevé dans les anomalies :
        0 attendu. C'est le contrôle de premier niveau de la consolidation."""
        for libelle, valeur, _ in self.anomalies:
            if libelle.startswith("Écart total transactions"):
                return float(valeur)
        return 0.0

    @property
    def sans_vente(self) -> bool:
        """Journée sans aucune vente (fermeture...). Cas normal, pas une
        anomalie de format : LightSpeed produit quand même ses rapports, vides.
        Signalé partout où le résultat est présenté, pour que la journée ne
        passe jamais inaperçue — sans faire échouer la consolidation."""
        return self.nb_tickets == 0

    @property
    def anomalies_a_verifier(self) -> list:
        """Anomalies méritant une vérification humaine, à l'exclusion des
        lignes purement informatives (cf. LIBELLES_INFORMATIFS)."""
        return [a for a in self.anomalies if not est_informatif(a[0])]

    @property
    def sans_anomalie_bloquante(self) -> bool:
        """Un écart non nul entre transactions et tickets est le seul cas où
        le classeur ne doit pas être considéré comme fiable. Les autres
        anomalies (groupes non mappés, périodes divergentes) sont des points
        à vérifier, pas des erreurs de calcul."""
        return abs(self.ecart_controle) <= 0.01


def famille(groupe: str) -> str:
    g = str(groupe)
    for pref in ("BEV", "FOOD", "DIV"):
        if g.startswith(pref):
            return pref
    for k, v in MAPPING_GROUPES.items():
        if g.startswith(k):
            return v
    return "AUTRE"


def periode(ts, periodes) -> str:
    """Période de service d'après l'heure d'ouverture (minutes depuis minuit)."""
    mn = ts.hour * 60 + ts.minute
    debuts = [d for _, _, d in periodes]
    for (nom, _, d), suivant in zip(periodes, debuts[1:] + [None]):
        if suivant is None or d <= mn < suivant:
            if mn >= d:
                return nom
    return periodes[-1][0]  # avant le premier début : nuit de la veille


def classer_fichiers(noms: list[str]) -> tuple[list[str], list[str], list[str]]:
    """Répartit des noms de fichiers en (tickets, transactions, non reconnus)
    d'après la convention de nommage des exports Lightspeed
    (`..._tickets_...xls` / `..._transactions_...xls`).

    Sert à pré-remplir le choix à l'écran quand plusieurs fichiers sont
    déposés d'un coup — jamais à décider seul : l'utilisateur garde la main,
    un export renommé à la main ne doit pas bloquer le traitement."""
    tickets, transactions, inconnus = [], [], []
    for nom in noms:
        bas = nom.lower()
        if re.search(r"_tickets?[_.]", bas):
            tickets.append(nom)
        elif re.search(r"_transactions?[_.]", bas):
            transactions.append(nom)
        else:
            inconnus.append(nom)
    return tickets, transactions, inconnus


# Mots-clés permettant de reconnaître le site dans le nom des exports : le
# nom de fichier porte la source ("..._barutopic_tickets_..."), mais pas
# toujours le mot "bar" — il peut se limiter à la marque de l'établissement.
# Même principe que MOTS_CLES_BAR page Convertisseur.
MOTS_CLES_SITE = {
    "BAR": ("bar", "utopic"),
    "RESTAURANT": ("restaurant", "aspp"),
}


def deviner_site(noms: list[str]) -> str | None:
    """Site suggéré d'après le nom des fichiers déposés, ou None si aucun
    indice ou si plusieurs sites sont évoqués — mieux vaut ne rien proposer
    qu'imposer un site arbitraire quand le nom est ambigu.

    Ne sert qu'à pré-remplir le choix à l'écran : le site reste sélectionné
    explicitement, et c'est cette valeur, jamais le nom de fichier, qui
    détermine les périodes de service appliquées au calcul."""
    trouves = {
        site for site, mots in MOTS_CLES_SITE.items()
        if any(mot in nom.lower() for nom in noms for mot in mots)
    }
    return trouves.pop() if len(trouves) == 1 else None


def _decoder(contenu: bytes, nom: str) -> str:
    """Même cascade d'encodages que le parser de la conversion comptable : les
    exports Lightspeed sortent tantôt en UTF-8 avec BOM, tantôt en cp1252."""
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return contenu.decode(enc)
        except UnicodeDecodeError:
            continue
    raise SyntheseError(f"« {nom} » : encodage du CSV non reconnu.")


def _nombre(v):
    """Texte d'un CSV -> nombre, en tolérant la virgule décimale et les espaces
    de milliers (y compris insécables). Une valeur vide vaut 0, comme une
    cellule vide d'un classeur. Reprend la logique de core.lightspeed_parser."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return 0.0
    if isinstance(v, str):
        v = v.replace("\xa0", "").replace(" ", "").replace(",", ".")
        if v in ("", "-"):
            return 0.0
    try:
        f = float(v)
    except (TypeError, ValueError):
        return 0.0
    return 0.0 if pd.isna(f) else f


def _lire_un(nom: str, contenu: bytes, libelle: str, colonnes_numeriques: list[str]) -> pd.DataFrame:
    """Un export, quel que soit son format. Le chemin .xls/.xlsx est laissé
    strictement inchangé : c'est celui vérifié contre le classeur de référence,
    et pandas y type déjà les colonnes."""
    if not nom.lower().endswith(".csv"):
        moteur = "xlrd" if nom.lower().endswith(".xls") else None
        return pd.read_excel(io.BytesIO(contenu), engine=moteur)

    texte = _decoder(contenu, nom)
    try:
        dialecte = csv.Sniffer().sniff("\n".join(texte.splitlines()[:5]), delimiters=";,\t")
        sep = dialecte.delimiter
    except csv.Error:
        premiere = texte.split("\n", 1)[0]
        sep = ";" if premiere.count(";") >= premiere.count(",") else ","
    df = pd.read_csv(io.StringIO(texte), sep=sep, dtype=str)
    for colonne in colonnes_numeriques:
        if colonne in df.columns:
            df[colonne] = df[colonne].map(_nombre)
    return df


def _lire(fichiers: list[tuple[str, bytes]], libelle: str, colonnes: list[str],
          colonnes_numeriques: list[str]) -> pd.DataFrame:
    """Concatène plusieurs exports du même rapport et dédoublonne par
    identifiant : permet de traiter un mois complet en déposant tous les
    fichiers d'un coup, sans compter deux fois un jour présent dans deux
    exports qui se chevauchent."""
    if not fichiers:
        raise SyntheseError(f"Aucun fichier « {libelle} » fourni.")
    frames = []
    for nom, contenu in fichiers:
        try:
            frames.append(_lire_un(nom, contenu, libelle, colonnes_numeriques))
        except SyntheseError:
            raise
        except Exception as e:
            raise SyntheseError(f"« {nom} » illisible comme export Lightspeed ({libelle}) : {e}") from e
    df = pd.concat(frames, ignore_index=True)
    manquantes = [c for c in colonnes if c not in df.columns]
    if manquantes:
        raise SyntheseError(
            f"Le rapport « {libelle} » n'a pas la structure attendue : colonne(s) "
            f"{', '.join(manquantes)} absente(s). Vérifiez qu'il ne s'agit pas de "
            f"l'autre rapport (Tickets et Transactions sont intervertis ?)."
        )
    return df.drop_duplicates(subset=["Identifier"])


def charger(tickets: list[tuple[str, bytes]], transactions: list[tuple[str, bytes]], periodes):
    """Lit les deux rapports et les croise. Renvoie (tickets, transactions,
    lignes fusionnées) — la fusion porte la période du TICKET sur chaque ligne
    de transaction, c'est elle qui alimente les totaux par période."""
    t = _lire(tickets, "Tickets", COLONNES_TICKETS, COLONNES_NUMERIQUES_TICKETS)
    x = _lire(transactions, "Transactions", COLONNES_TRANSACTIONS, COLONNES_NUMERIQUES_TRANSACTIONS)
    t["Date"] = pd.to_datetime(t["Date"], format="%d/%m/%y %H:%M")
    t["OpenDate"] = pd.to_datetime(t["OpenDate"], format="%d/%m/%y %H:%M")
    t["Jour"] = (t["OpenDate"] - DEBUT_JOURNEE).dt.date
    t["ProfilLightspeed"] = t["Profil"]
    t["Profil"] = t["OpenDate"].map(lambda ts: periode(ts, periodes))      # période à l'OUVERTURE (référence)
    t["ProfilReglement"] = t["Date"].map(lambda ts: periode(ts, periodes))  # période au RÈGLEMENT (déclinaison)
    # Table physique : "BAR, Table 35" / "Table 3.1" (split) -> "35" / "3"
    t["Table"] = pd.to_numeric(t["AccountName"].astype(str).str.extract(r"Table\s*(\d+)")[0]).astype("Int64")
    # Une "rotation" = une table ouverte pour un service : ticket SALE ou RECALL
    # non annulé, sur une table identifiée. Les SPLIT (additions séparées) et
    # les VOID ne créent pas d'ouverture supplémentaire.
    t["Rotation"] = (t["Type"].isin(["SALE", "RECALL"]) & (t["Annulée"] != "Oui")
                     & t["Table"].notna()).astype(int)
    # Durée de présence (minutes) entre ouverture et règlement, sur les rotations uniquement
    duree = (t["Date"] - t["OpenDate"]).dt.total_seconds() / 60
    t["DureeMin"] = duree.where(t["Rotation"] == 1).round(0)
    tk = t[["Account", "Profil", "ProfilReglement", "Type", "Jour", "Couverts", "Total", "PreTax", "Annulée"]]
    tk = tk.rename(columns={"Profil": "ProfilTicket", "Type": "TypeTicket",
                            "Total": "TotalTicket", "PreTax": "PreTaxTicket"})
    m = x.merge(tk.drop(columns=["Couverts", "TotalTicket", "PreTaxTicket", "Annulée"]),
                on="Account", how="left")
    m["Famille"] = m["Group"].map(famille)
    return t, x, m


def anomalies(t, x, m, periodes) -> list[tuple]:
    a = []
    if t.empty:
        a.append((
            "Aucune vente sur la période", 0,
            "Le rapport Tickets ne contient aucune ligne : journée sans activité "
            "(fermeture, jour férié...). Tous les totaux sont à zéro. Ce n'est pas "
            "une anomalie de format — mais si la journée aurait dû être ouverte, "
            "c'est l'export LightSpeed qu'il faut vérifier.",
        ))
    orphelins = m[m["ProfilTicket"].isna()]
    if len(orphelins):
        a.append(("Lignes de transaction sans ticket", len(orphelins),
                  ", ".join(sorted(set(orphelins["Account"].astype(str)))[:20])))
    autres = m[m["Famille"] == "AUTRE"]
    if len(autres):
        a.append(("Groupes non mappés (famille AUTRE)", len(autres),
                  ", ".join(sorted(set(autres["Group"].astype(str))))))
    ecart = round(x["FinalPrice"].sum() - t["Total"].sum(), 2)
    a.append(("Écart total transactions - tickets (TTC)", ecart, "0 attendu"))
    for p, g in m.groupby("ProfilTicket"):
        e = round(g["FinalPrice"].sum() - t[t["Profil"] == p]["Total"].sum(), 2)
        if abs(e) > 0.01:
            a.append((f"Écart par période : {p}", e, "0 attendu"))
    hors = t[~t["Profil"].isin([p for p, *_ in periodes])]
    if len(hors):
        a.append(("Tickets avec un profil hors liste", len(hors),
                  ", ".join(sorted(set(hors["Profil"].astype(str))))))
    diff = t[(t["Profil"] != t["ProfilLightspeed"]) & (t["Type"] != "VOID")]
    if len(diff):
        a.append(("Tickets rattachés à leur période d'ouverture, et non au profil Lightspeed de clôture", len(diff),
                  ", ".join(f"{i}: {a_} -> {b_}" for i, a_, b_ in
                            zip(diff["Identifier"], diff["ProfilLightspeed"], diff["Profil"]))))
    annules = t[t["Annulée"] == "Oui"]
    if len(annules):
        a.append(("Tickets annulés (VOID, comptés en négatif)", len(annules),
                  ", ".join(annules["Identifier"].astype(str))))
    return a


def styliser_entete(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, size=9, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=TAUPE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _jour_depuis_ddmmaa(valeur: str | None):
    """Convertit une date "dd/mm/aa" (format des noms d'export, cf.
    core.email_ingest) en date. None si elle n'est pas exploitable."""
    try:
        j, mo, an = str(valeur).split("/")
        return dt.date(2000 + int(an), int(mo), int(j))
    except (AttributeError, ValueError):
        return None


def ecrire(t, x, m, periodes, titre_site: str, jours=None) -> bytes:
    """Construit le classeur et le renvoie en octets : aucune écriture disque,
    l'app le propose au téléchargement et l'archive elle-même."""
    wb = Workbook()
    fin = Side(style="thin", color=GRIS_CLAIR)
    bord = Border(bottom=fin)
    eur = '#,##0.00 "€";-#,##0.00 "€";-'
    pct = "0.0%;-0.0%;-"
    entier = "#,##0;-#,##0;-"

    # ---------- DONNEES : une ligne par transaction, prête pour SUMIFS ----------
    wd = wb.active
    wd.title = "DONNEES"
    cols = ["Jour", "ProfilTicket", "Account", "Type", "Item", "Group", "Famille",
            "Qty", "FinalPrice", "PreTax", "TaxAmount", "TaxName", "ProfilReglement"]
    wd.append(cols)
    styliser_entete(wd, 1, 1, len(cols))
    for r in m[cols].itertuples(index=False):
        wd.append([v if not (isinstance(v, float) and pd.isna(v)) else None for v in r])
    for c in range(1, len(cols) + 1):
        wd.column_dimensions[get_column_letter(c)].width = 16
    wd.freeze_panes = "A2"
    n = len(m) + 1

    # ---------- COUVERTS : une ligne par ticket ----------
    wc = wb.create_sheet("TICKETS")
    tc = ["Jour", "Profil", "Identifier", "Account", "Type", "Annulée", "Couverts", "Total", "PreTax",
          "Table", "Rotation", "OpenDate", "Date", "ProfilLightspeed", "ProfilReglement", "DureeMin"]
    wc.append(tc)
    styliser_entete(wc, 1, 1, len(tc))
    for r in t[tc].itertuples(index=False):
        wc.append([None if (v is pd.NA or (isinstance(v, float) and pd.isna(v))) else (int(v) if isinstance(v, (pd.Int64Dtype, )) else v) for v in r])
    for c in range(1, len(tc) + 1):
        wc.column_dimensions[get_column_letter(c)].width = 16
    for r in range(2, len(t) + 2):
        wc.cell(r, 12).number_format = wc.cell(r, 13).number_format = "DD/MM/YY HH:MM"
    nt = len(t) + 1

    # ---------- SYNTHESE : mise en forme du modèle DAF ----------
    ws = wb.create_sheet("SYNTHESE", 0)
    # Fournis par l'appelant : une journée sans aucune vente n'en contient
    # aucun, et le classeur doit tout de même savoir de quelle période il
    # parle — elle vient alors du nom de fichier.
    if jours is None:
        jours = sorted(set(t["Jour"]))
    D, T = "DONNEES", "TICKETS"
    FN = "Aptos Narrow"
    med, thin = Side(style="medium"), Side(style="thin")
    eur_daf = '_ * #,##0.00_)\\ "€"_ ;_ * \\(#,##0.00\\)\\ "€"_ ;_ * "-"??_)\\ "€"_ ;_ @_ '
    ent_daf = '_-* #,##0_-;\\-* #,##0_-;_-* "-"??_-;_-@_-'
    titre = titre_site

    def bloc(r0, libelle, col_profil_d, col_profil_t):
        """Écrit un bloc de synthèse à partir de la ligne r0. Renvoie la ligne suivante libre."""
        ws.cell(r0, 3, titre).font = Font(name="Arial", size=12, bold=True)
        ws.cell(r0, 3).alignment = Alignment(horizontal="center")
        ws.cell(r0, 5, libelle).font = Font(name=FN, size=10, italic=True, color=GRIS_MOYEN)
        ws.row_dimensions[r0].height = 17
        hdr = ["CA TTC", "CA HT", "% Du total", "Couverts", "CA Food HT", "CA BEV HT", "Tables ouvertes"]
        rh = r0 + 1
        for i, h in enumerate(hdr):
            c = ws.cell(rh, 3 + i, h)
            c.font = Font(name=FN, size=11, bold=True)
            c.alignment = Alignment(horizontal="center")
            c.border = Border(top=med, right=med if i == len(hdr) - 1 else None)
        rp = rh + 1
        for i, (p, h, _) in enumerate(periodes):
            r = rp + i
            ws.cell(r, 1, p).border = Border(left=med)
            ws.cell(r, 2, h)
            ws.cell(r, 3, f"=SUMIFS({D}!$I$2:$I${n},{D}!${col_profil_d}$2:${col_profil_d}${n},$A{r})")
            ws.cell(r, 4, f"=SUMIFS({D}!$J$2:$J${n},{D}!${col_profil_d}$2:${col_profil_d}${n},$A{r})")
            ws.cell(r, 5, f"=IF($C${rp+len(periodes)}=0,0,C{r}/$C${rp+len(periodes)})")
            ws.cell(r, 6, f"=SUMIFS({T}!$G$2:$G${nt},{T}!${col_profil_t}$2:${col_profil_t}${nt},$A{r})")
            ws.cell(r, 7, f'=SUMIFS({D}!$J$2:$J${n},{D}!${col_profil_d}$2:${col_profil_d}${n},$A{r},{D}!$G$2:$G${n},"FOOD")')
            ws.cell(r, 8, f'=SUMIFS({D}!$J$2:$J${n},{D}!${col_profil_d}$2:${col_profil_d}${n},$A{r},{D}!$G$2:$G${n},"BEV")')
            ws.cell(r, 9, f"=SUMIFS({T}!$K$2:$K${nt},{T}!${col_profil_t}$2:${col_profil_t}${nt},$A{r})")
            ws.cell(r, 9).border = Border(right=med)
            for c in range(1, 10):
                ws.cell(r, c).font = Font(name=FN, size=12)
            for c in (3, 4, 7, 8):
                ws.cell(r, c).number_format = eur_daf
            ws.cell(r, 5).number_format = "0.0%"
            ws.cell(r, 6).number_format = ent_daf
            ws.cell(r, 9).number_format = ent_daf
        rt = rp + len(periodes)
        ws.cell(rt, 1, "Total des ventes").font = Font(name=FN, size=12)
        ws.cell(rt, 1).border = Border(left=med, bottom=med)
        ws.cell(rt, 2).border = Border(bottom=med)
        for c in range(3, 10):
            L = get_column_letter(c)
            cell = ws.cell(rt, c, f"=SUM({L}{rp}:{L}{rt-1})" if c != 5 else f"=IF(C{rt}=0,0,SUM(E{rp}:E{rt-1}))")
            cell.font = Font(name=FN, size=11, bold=True)
            cell.border = Border(top=thin, bottom=med, right=med if c == 9 else None)
            cell.number_format = eur_daf if c in (3, 4, 7, 8) else ("0.0%" if c == 5 else ent_daf)
        ws.row_dimensions[rt].height = 17
        return rt + 3

    ws["A1"] = f"Du {jours[0]:%d/%m/%Y} au {jours[-1]:%d/%m/%Y} — {len(jours)} jour(s) d'exploitation"
    ws["A1"].font = Font(name=FN, size=10, color=GRIS_MOYEN)
    nxt = bloc(3, "Période = heure d'OUVERTURE de la table", "B", "B")
    nxt = bloc(nxt, "Période = heure de RÈGLEMENT de la table", "M", "O")
    ws.cell(nxt, 1, "Contrôle : total DONNEES - total TICKETS (0 attendu)").font = Font(name=FN, size=9, color=GRIS_MOYEN)
    ws.cell(nxt, 3, f"=SUM({D}!I2:I{n})-SUM({T}!H2:H{nt})").number_format = eur_daf
    ws.cell(nxt + 1, 1, "CA Food HT + CA BEV HT + Divers HT = CA HT. CA = somme de toutes les lignes de transaction du ticket, tous types confondus. Tables ouvertes = tickets SALE/RECALL non annulés sur une table identifiée (les additions séparées ne comptent pas).").font = Font(name=FN, size=9, color=GRIS_MOYEN)
    for L, w in zip("ABCDEFGHI", (17, 15, 21.8, 15, 18.3, 23.5, 16.5, 19.5, 16)):
        ws.column_dimensions[L].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------- JOUR x PERIODE ----------
    wj = wb.create_sheet("JOUR x PERIODE", 1)
    wj["A1"] = "CA TTC ET COUVERTS PAR JOUR ET PÉRIODE (période à l'ouverture)"
    wj["A1"].font = Font(name=FONT, size=14)
    h2 = ["Jour"]
    for p, *_ in periodes:
        h2 += [f"{p} TTC", f"{p} couverts"]
    h2 += ["Total TTC", "Total couverts"]
    wj.append([])
    wj.append(h2)
    styliser_entete(wj, 3, 1, len(h2))
    for k, j in enumerate(jours):
        r = 4 + k
        wj.cell(r, 1, j).number_format = "DD/MM/YYYY"
        for i, (p, *_) in enumerate(periodes):
            cca, ccv = 2 + 2 * i, 3 + 2 * i
            wj.cell(r, cca, f'=SUMIFS({D}!$I$2:$I${n},{D}!$B$2:$B${n},"{p}",{D}!$A$2:$A${n},$A{r})').number_format = eur
            wj.cell(r, ccv, f'=SUMIFS({T}!$G$2:$G${nt},{T}!$B$2:$B${nt},"{p}",{T}!$A$2:$A${nt},$A{r})').number_format = entier
        ctot = 2 + 2 * len(periodes)
        ca_cols = ",".join(f"{get_column_letter(2+2*i)}{r}" for i in range(len(periodes)))
        cv_cols = ",".join(f"{get_column_letter(3+2*i)}{r}" for i in range(len(periodes)))
        wj.cell(r, ctot, f"=SUM({ca_cols})").number_format = eur
        wj.cell(r, ctot + 1, f"=SUM({cv_cols})").number_format = entier
        for c in range(1, len(h2) + 1):
            wj.cell(r, c).font = Font(name=FONT, size=10, color=ANTHRACITE)
            wj.cell(r, c).border = bord
    for c in range(1, len(h2) + 1):
        wj.column_dimensions[get_column_letter(c)].width = 14
    wj.freeze_panes = "B4"

    # ---------- ROTATIONS : tables x période ----------
    wr = wb.create_sheet("ROTATIONS", 2)
    wr["A1"] = "OUVERTURES PAR TABLE ET PÉRIODE (période à l'ouverture, toute la durée du rapport)"
    wr["A1"].font = Font(name=FONT, size=14)
    tables = sorted({int(v) for v in t["Table"].dropna()})
    h3 = ["Table"] + [p for p, *_ in periodes] + ["Total ouvertures", "Ouvertures / jour", "Couverts"]
    wr.append([])
    wr.append(h3)
    styliser_entete(wr, 3, 1, len(h3))
    for k, tb in enumerate(tables):
        r = 4 + k
        wr.cell(r, 1, tb)
        for i, (p, *_) in enumerate(periodes):
            wr.cell(r, 2 + i, f'=SUMIFS({T}!$K$2:$K${nt},{T}!$J$2:$J${nt},$A{r},{T}!$B$2:$B${nt},"{p}")').number_format = entier
        ct = 2 + len(periodes)
        wr.cell(r, ct, f"=SUM(B{r}:{get_column_letter(ct-1)}{r})").number_format = entier
        wr.cell(r, ct + 1, f"={get_column_letter(ct)}{r}/{len(jours)}").number_format = "0.0"
        wr.cell(r, ct + 2, f"=SUMIFS({T}!$G$2:$G${nt},{T}!$J$2:$J${nt},$A{r})").number_format = entier
        for c in range(1, len(h3) + 1):
            wr.cell(r, c).font = Font(name=FONT, size=10, color=ANTHRACITE)
            wr.cell(r, c).border = bord
    for c in range(1, len(h3) + 1):
        wr.column_dimensions[get_column_letter(c)].width = 15

    # ---------- DUREE PRESENCE ----------
    wp = wb.create_sheet("DUREE PRESENCE", 3)
    wp["A1"] = "DURÉE DE PRÉSENCE DES CLIENTS (ouverture → règlement, par période d'ouverture)"
    wp["A1"].font = Font(name=FONT, size=14)
    wp["A2"] = "Base : tables ouvertes (tickets SALE/RECALL non annulés). Les additions séparées ne comptent pas."
    wp["A2"].font = Font(name=FONT, size=8, color=GRIS_MOYEN)
    tranches = [("< 30 min", 0, 30), ("30 min – 1h", 30, 60), ("1h – 1h30", 60, 90),
                ("1h30 – 2h", 90, 120), ("2h – 2h30", 120, 150), ("2h30 – 3h", 150, 180),
                ("> 3h", 180, None)]
    noms = [p for p, *_ in periodes]
    h4 = ["Tranche"] + noms + ["Total", "% des tables"]
    wp.append([])
    wp.append(h4)
    styliser_entete(wp, 4, 1, len(h4))
    P = f"{T}!$P$2:$P${nt}"
    PR = f"{T}!$B$2:$B${nt}"
    r1 = 5
    for k, (lib, lo, hi) in enumerate(tranches):
        r = r1 + k
        wp.cell(r, 1, lib)
        for i, p in enumerate(noms):
            crit = f'{P},">="&{lo}' + (f',{P},"<"&{hi}' if hi is not None else "")
            wp.cell(r, 2 + i, f'=COUNTIFS({PR},"{p}",{crit})').number_format = entier
        ct = 2 + len(noms)
        wp.cell(r, ct, f"=SUM(B{r}:{get_column_letter(ct-1)}{r})").number_format = entier
        wp.cell(r, ct + 1, f"=IF({get_column_letter(ct)}${r1+len(tranches)}=0,0,{get_column_letter(ct)}{r}/{get_column_letter(ct)}${r1+len(tranches)})").number_format = pct
    rt = r1 + len(tranches)
    wp.cell(rt, 1, "Total tables")
    for c in range(2, 3 + len(noms)):
        L = get_column_letter(c)
        wp.cell(rt, c, f"=SUM({L}{r1}:{L}{rt-1})").number_format = entier
    wp.cell(rt, 3 + len(noms), f"=IF({get_column_letter(2+len(noms))}{rt}=0,0,1)").number_format = pct
    # Statistiques
    rs = rt + 2
    wp.cell(rs, 1, "Statistiques (minutes)")
    styliser_entete(wp, rs, 1, len(h4) - 1)
    for i, p in enumerate(noms):
        wp.cell(rs, 2 + i, p)
    wp.cell(rs, 2 + len(noms), "Toutes")
    stats = [("Durée moyenne", "AVERAGEIFS", True), ("Durée minimum", "_xlfn.MINIFS", False),
             ("Durée maximum", "_xlfn.MAXIFS", False)]
    for k, (lib, fn, _) in enumerate(stats):
        r = rs + 1 + k
        wp.cell(r, 1, lib)
        for i, p in enumerate(noms):
            wp.cell(r, 2 + i, f'=IFERROR({fn}({P},{PR},"{p}",{P},">=0"),0)').number_format = "0"
        allfn = {"AVERAGEIFS": "AVERAGE", "_xlfn.MINIFS": "MIN", "_xlfn.MAXIFS": "MAX"}[fn]
        wp.cell(r, 2 + len(noms), f"=IFERROR({allfn}({P}),0)").number_format = "0"
    r = rs + 1 + len(stats)
    wp.cell(r, 1, "Durée moyenne (h:mm)")
    for c in range(2, 3 + len(noms)):
        L = get_column_letter(c)
        wp.cell(r, c, f"={L}{rs+1}/1440").number_format = "[h]:mm"
    for row in wp.iter_rows(min_row=r1, max_row=r):
        for cell in row:
            if cell.row not in (rs,):
                cell.font = Font(name=FONT, size=10, color=ANTHRACITE)
                cell.border = bord
    for c in range(1, len(h4) + 1):
        wp.column_dimensions[get_column_letter(c)].width = 16
    wp.column_dimensions["A"].width = 24

    # ---------- ANOMALIES ----------
    wa = wb.create_sheet("ANOMALIES")
    wa.append(["Contrôle", "Valeur", "Détail"])
    styliser_entete(wa, 1, 1, 3)
    for row in anomalies(t, x, m, periodes):
        wa.append(list(row))
    for c, w in zip("ABC", (45, 12, 80)):
        wa.column_dimensions[c].width = w
    for row in wa.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10, color=ANTHRACITE)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def construire_synthese(
    tickets: list[tuple[str, bytes]],
    transactions: list[tuple[str, bytes]],
    site: str,
    periode: tuple[str | None, str | None] | None = None,
) -> SyntheseResult:
    """Point d'entrée unique de la consolidation : deux jeux d'exports en
    octets (nom, contenu) et un site, un classeur en octets et de quoi le
    présenter. Plusieurs fichiers par rapport sont acceptés pour couvrir une
    période plus longue qu'une journée (dédoublonnage par identifiant).

    `periode` : (date de début, date de fin) au format "dd/mm/aa", telles que
    lues dans le nom des exports (cf. core.email_ingest.extraire_periode).
    Servent uniquement quand les rapports ne contiennent AUCUNE vente : le
    classeur doit alors savoir de quelle journée il parle, et le contenu ne
    peut plus le lui dire. Une journée sans vente est un cas normal - le
    restaurant était fermé - et produit un classeur à zéro, signalé comme tel,
    plutôt qu'un échec : même décision que pour la conversion comptable, où
    une journée sans vente a cessé d'être une erreur."""
    conf = SITES.get(str(site).upper())
    if conf is None:
        raise SyntheseError(
            f"Site « {site} » inconnu — attendu : {', '.join(SITES)}."
        )
    periodes = conf["periodes"]
    t, x, m = charger(tickets, transactions, periodes)

    res = SyntheseResult(site=str(site).upper())
    res.fichiers_sources = [nom for nom, _ in tickets] + [nom for nom, _ in transactions]
    res.jours = sorted(set(t["Jour"]))
    if not res.jours:
        # Aucune vente : la période vient du nom de fichier. À défaut, le
        # classeur serait sans date - on préfère un repli sur le jour même à
        # un plantage ou à un refus de traiter.
        debut = _jour_depuis_ddmmaa((periode or (None, None))[0])
        res.jours = [debut or dt.date.today()]
    res.nb_tickets = len(t)
    res.nb_lignes = len(m)
    # Mêmes colonnes que celles totalisées par la feuille SYNTHESE (DONNEES!I
    # et DONNEES!J) : les chiffres affichés à l'écran et ceux du classeur ne
    # peuvent donc pas diverger.
    res.ca_ttc = round(float(m["FinalPrice"].sum()), 2)
    res.ca_ht = round(float(m["PreTax"].sum()), 2)
    res.couverts = int(t["Couverts"].sum())
    res.anomalies = anomalies(t, x, m, periodes)
    res.classeur = ecrire(t, x, m, periodes, conf["titre"], res.jours)
    return res
